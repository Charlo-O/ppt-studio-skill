#!/usr/bin/env python3
"""Export a deck to PowerPoint and PDF.

  export_pptx.py DECK [--editable] [--images] [--pdf]
                 [--name NAME] [--slides s01,s02] [--windows-fonts]
                 [--font-map "PingFang SC=Microsoft YaHei" ...] [--keep-work]

With no format flag: editable-mode decks get all three outputs; image-mode decks get the image
PPTX and PDF built from the same Agent-authored HTML renders in out/png/.

  out/<name>.pptx         editable: live text boxes, native shapes, every image its own object
  out/<name>-images.pptx  one full-bleed picture per slide (pixel-exact, text not editable)
  out/<name>.pdf          one page per slide
Speaker notes come from deck.json `slides[].notes`. Opaque images are stored as high-quality JPEG
to keep files small; --lossless keeps PNG everywhere.
"""

from __future__ import annotations

import argparse
import importlib.util
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import SCRIPT_DIR, canvas_size, die, find_deck, load_deck, slide_files  # noqa: E402

WINDOWS_FONTS = {
    "PingFang SC": "Microsoft YaHei", "PingFang TC": "Microsoft JhengHei", "Hiragino Sans GB": "Microsoft YaHei",
    "STHeiti": "Microsoft YaHei", "Heiti SC": "Microsoft YaHei", "Songti SC": "SimSun", "STSong": "SimSun",
    "Kaiti SC": "KaiTi", "STKaiti": "KaiTi", "Yuanti SC": "Microsoft YaHei", "Hiragino Sans": "Yu Gothic",
    "Hiragino Kaku Gothic ProN": "Yu Gothic", "Hiragino Mincho ProN": "Yu Mincho",
    "Helvetica Neue": "Arial", "Helvetica": "Arial", "Avenir Next": "Segoe UI", "Avenir": "Segoe UI",
    "Avenir Next Condensed": "Arial Narrow", "Futura": "Century Gothic", "Gill Sans": "Gill Sans MT",
    "Didot": "Bodoni MT", "Bodoni 72": "Bodoni MT", "Baskerville": "Baskerville Old Face",
    "Menlo": "Consolas", "SF Mono": "Consolas", "DIN Alternate": "Bahnschrift",
    "DIN Condensed": "Bahnschrift Condensed", "American Typewriter": "Courier New",
}


def load_converter():
    path = SCRIPT_DIR / "vendor" / "html_to_pptx.py"
    spec = importlib.util.spec_from_file_location("ppt_html_to_pptx", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def deck_name(deck: Path, data: dict, explicit: str | None) -> str:
    raw = explicit or data.get("title") or deck.name
    name = re.sub(r"[\\/:*?\"<>|\s]+", "-", raw).strip("-.")
    return name or "deck"


def notes_by_id(data: dict) -> dict:
    return {s.get("id"): (s.get("notes") or "").strip() for s in data.get("slides") or []}


def set_notes(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # group
            yield from iter_shapes(shape.shapes)


def remap_fonts(prs, mapping: dict) -> dict:
    from pptx.oxml.ns import qn

    used: dict[str, int] = {}
    for slide in prs.slides:
        for shape in iter_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    rpr = run._r.get_or_add_rPr()
                    for tag in ("a:latin", "a:ea", "a:cs"):
                        node = rpr.find(qn(tag))
                        if node is None:
                            continue
                        face = node.get("typeface")
                        if face in mapping:
                            node.set("typeface", mapping[face])
                            face = mapping[face]
                        if face and tag in ("a:latin", "a:ea"):
                            used[face] = used.get(face, 0) + 1
    return used


def slim_scene(scene: dict, quality: int = 90) -> None:
    """Store opaque rasters as JPEG; only images with real transparency stay PNG."""
    from PIL import Image

    def to_jpeg(path_str: str) -> str:
        path = Path(path_str)
        with Image.open(path) as im:
            if "A" in im.getbands():
                alpha = im.getchannel("A")
                if alpha.getextrema()[0] < 250:
                    return path_str
            target = path.with_suffix(".jpg")
            im.convert("RGB").save(target, quality=quality, optimize=True, progressive=True)
        return str(target)

    scene["background_png"] = to_jpeg(scene["background_png"])
    for element in scene["elements"]:
        if element.get("partPng"):
            element["partPng"] = to_jpeg(element["partPng"])


def export_editable(deck: Path, files: list[Path], out: Path, notes: dict, mapping: dict, keep: bool,
                    lossless: bool = False) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    conv = load_converter()
    work = Path(tempfile.mkdtemp(prefix="pptx-work-", dir=deck / "out"))
    prs = Presentation()
    summary = []
    try:
        for index, path in enumerate(files):
            scene = conv.capture_scene(path, work / path.stem, serve_root=deck, scale=2)
            if not lossless:
                slim_scene(scene)
            if index == 0:
                width, height = conv.slide_size_emu(scene["canvas"]["width"], scene["canvas"]["height"])
                prs.slide_width, prs.slide_height = Emu(width), Emu(height)
            counts = conv.build_pptx(scene, None, prs=prs)
            set_notes(prs.slides[-1], notes.get(path.stem, ""))
            summary.append((path.stem, counts))
            print(f"  {path.stem}: {counts['text_count']} text boxes · {counts['shape_count']} shapes · "
                  f"{counts['picture_count']} pictures")
        fonts = remap_fonts(prs, mapping)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
    finally:
        if not keep:
            shutil.rmtree(work, ignore_errors=True)
    check = Presentation(str(out))
    if len(check.slides) != len(files):
        die(f"editable deck has {len(check.slides)} slides, expected {len(files)}")
    print(f"✓ editable PPTX → {out}  ({len(files)} slides, {out.stat().st_size / 1e6:.1f} MB)")
    if fonts:
        print("  fonts referenced: " + ", ".join(f"{k} ({v})" for k, v in sorted(fonts.items(), key=lambda kv: -kv[1])))


def fitted_images(paths: list[Path], size: tuple[int, int], work: Path, lossless: bool = False) -> list[Path]:
    """Slide images at the exact canvas size; JPEG (quality 92) unless lossless is asked for."""
    from PIL import Image

    from imgtool import fit_image

    work.mkdir(parents=True, exist_ok=True)
    result = []
    for path in paths:
        with Image.open(path) as im:
            rgb = im.convert("RGB")
            if rgb.size != size:
                rgb = fit_image(rgb, size[0], size[1])
            target = work / f"{path.stem}.{'png' if lossless else 'jpg'}"
            if lossless:
                rgb.save(target)
            else:
                rgb.save(target, quality=92, optimize=True, progressive=True)
            result.append(target)
    return result


def export_images(images: list[Path], ids: list[str], out: Path, notes: dict, canvas: tuple[int, int]) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    conv = load_converter()
    prs = Presentation()
    width, height = conv.slide_size_emu(*canvas)
    prs.slide_width, prs.slide_height = Emu(width), Emu(height)
    for image, slide_id in zip(images, ids):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), Emu(0), Emu(0), width=Emu(width), height=Emu(height))
        set_notes(slide, notes.get(slide_id, ""))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"✓ image PPTX → {out}  ({len(images)} slides, {out.stat().st_size / 1e6:.1f} MB)")


def export_pdf(images: list[Path], out: Path) -> None:
    from PIL import Image

    pages = [Image.open(p).convert("RGB") for p in images]
    out.parent.mkdir(parents=True, exist_ok=True)
    pages[0].save(out, save_all=True, append_images=pages[1:], resolution=144)
    print(f"✓ PDF → {out}  ({len(pages)} pages, {out.stat().st_size / 1e6:.1f} MB)")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export a deck to PPTX and PDF")
    parser.add_argument("deck")
    parser.add_argument("--editable", action="store_true")
    parser.add_argument("--images", action="store_true")
    parser.add_argument("--pdf", action="store_true")
    parser.add_argument("--name")
    parser.add_argument("--slides", help="comma-separated slide ids")
    parser.add_argument("--windows-fonts", action="store_true", help="swap macOS fonts for Windows equivalents")
    parser.add_argument("--font-map", action="append", default=[], help='"From Family=To Family"')
    parser.add_argument("--keep-work", action="store_true")
    parser.add_argument("--lossless", action="store_true", help="keep every raster as PNG (much larger files)")
    args = parser.parse_args()

    deck = find_deck(Path(args.deck))
    data = load_deck(deck)
    mode = data.get("mode", "editable")
    if not (args.editable or args.images or args.pdf):
        args.editable = mode == "editable"
        args.images = args.pdf = True
    name = deck_name(deck, data, args.name)
    notes = notes_by_id(data)
    canvas = canvas_size(deck)
    only = args.slides.split(",") if args.slides else None
    mapping = dict(WINDOWS_FONTS) if args.windows_fonts else {}
    for pair in args.font_map:
        if "=" not in pair:
            die(f"--font-map expects From=To, got {pair}")
        src, dst = pair.split("=", 1)
        mapping[src.strip()] = dst.strip()

    if args.editable:
        files = slide_files(deck, only)
        if not files:
            die("no slides/*.html to export; build the Agent-designed slides first")
        export_editable(deck, files, deck / "out" / f"{name}.pptx", notes, mapping, args.keep_work, args.lossless)

    if args.images or args.pdf:
        files = slide_files(deck, only)
        if not files:
            die("no slides/*.html to export; build the Agent-designed slides first")
        ids = [p.stem for p in files]
        pngs = [deck / "out" / "png" / f"{i}.png" for i in ids]
        shared = [deck / "theme.css", *(deck / "assets").rglob("*")]
        newest_shared = max((f.stat().st_mtime for f in shared if f.is_file()), default=0)
        stale = [p for p, f in zip(pngs, files) if not p.is_file()
                 or p.stat().st_mtime < max(f.stat().st_mtime, newest_shared)]
        if stale:
            print("rendering slides first …")
            subprocess.run([sys.executable, str(SCRIPT_DIR / "render.py"), str(deck)]
                           + (["--slides", args.slides] if args.slides else []), check=True)
        missing = [str(p) for p in pngs if not p.is_file()]
        if missing:
            die(f"missing slide images: {missing}")
        work = deck / "out" / ".fit"
        images = fitted_images(pngs, canvas, work, args.lossless)
        if args.images:
            export_images(images, ids, deck / "out" / f"{name}-images.pptx", notes, canvas)
        if args.pdf:
            export_pdf(images, deck / "out" / f"{name}.pdf")
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main()
